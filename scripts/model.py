"""Расчётная модель кейса: RFM-перераспределение бюджета.

Единственный источник цифр. Все вводные — допущения (в кейсе нет данных),
они собраны в одном месте и помечены. Скрипт пишет model/model.json,
model/rfm_budget_model.xlsx (если установлен openpyxl) и model/segments.csv.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "model"

# ---------------------------------------------------------------- допущения
# Месячные показатели, ₸. База: 50 000 клиентов с покупкой за 12 месяцев.
GROSS_MARGIN = 0.30          # валовая маржа e-commerce
NEW_CUSTOMERS = 8_000        # новые клиенты в месяц (приток стабилен)
NEW_AOV = 14_400             # средний чек первой покупки
SCENARIO_SHARE = {"pessimistic": 0.5, "base": 1.0, "optimistic": 1.3}

# RFM-сегменты: правило, клиенты, частота (заказов на клиента в месяц),
# средний чек, бюджет до/после (млн ₸/мес), целевые частота/чек.
SEGMENTS = [
    {
        "id": "champions", "name": "Чемпионы", "rule": "R 4–5 · F 4–5",
        "who": "Покупают часто и недавно, дают треть выручки базы",
        "customers": 4_000, "freq": 0.90, "aov": 18_000,
        "budget_before": 4.5, "budget_after": 1.5,
        "freq_after": 0.88, "aov_after": 18_000, "extra_revenue": 2.8,
        "goal": "Удержать без скидок, превратить в амбассадоров",
        "lever": "Снять общие скидки (купят и без них), дать нематериальные "
                 "привилегии и реферальную программу",
        "channels": "Push и e-mail в приложении, ранний доступ, закрытые "
                    "распродажи, реферальная ссылка «приведи друга»",
        "extra_note": "реферальная программа: 4 000 × 5% приглашений × 14 000 ₸",
    },
    {
        "id": "loyal", "name": "Лояльные", "rule": "R 3–5 · F 3–5",
        "who": "Регулярные покупатели среднего чека",
        "customers": 7_500, "freq": 0.45, "aov": 16_000,
        "budget_before": 3.0, "budget_after": 2.0,
        "freq_after": 0.47, "aov_after": 16_800, "extra_revenue": 0.0,
        "goal": "Рост чека и частоты через персональные рекомендации",
        "lever": "Кросс-селл и наборы по истории покупок, порог бесплатной "
                 "доставки выше текущего среднего чека",
        "channels": "Персональные подборки в e-mail и push, блок "
                    "«вам подойдёт» на сайте, программа баллов",
        "extra_note": "",
    },
    {
        "id": "potential", "name": "Потенциально лояльные", "rule": "R 4–5 · F 1–2",
        "who": "Новые клиенты после 1–2 покупок — здесь теряется повтор",
        "customers": 12_500, "freq": 0.25, "aov": 14_000,
        "budget_before": 2.5, "budget_after": 5.0,
        "freq_after": 0.35, "aov_after": 14_000, "extra_revenue": 0.0,
        "goal": "Довести до второй и третьей покупки за 60 дней",
        "lever": "Триггерная цепочка онбординга D+3 / D+14 / D+30, бонус на "
                 "вторую покупку, установка приложения",
        "channels": "E-mail и push-цепочка, напоминание о пополнении "
                    "расходников, бонус в приложении, ретаргет на "
                    "просмотренные товары",
        "extra_note": "",
    },
    {
        "id": "at_risk", "name": "Под угрозой оттока", "rule": "R 1–2 · F 3–5",
        "who": "Раньше покупали часто, но давно не возвращались",
        "customers": 9_000, "freq": 0.12, "aov": 15_000,
        "budget_before": 1.5, "budget_after": 3.5,
        "freq_after": 0.18, "aov_after": 15_000, "extra_revenue": 0.0,
        "goal": "Вернуть ценных клиентов, пока не ушли насовсем",
        "lever": "Win-back с персональным предложением по любимой категории, "
                 "опрос «почему перестали покупать», звонок топ-клиентам",
        "channels": "Персональный e-mail и SMS, ретаргетинг, звонок "
                    "колл-центра для клиентов с M 4–5",
        "extra_note": "",
    },
    {
        "id": "hibernating", "name": "Спящие и потерянные", "rule": "R 1–2 · F 1–2",
        "who": "Одна-две покупки давно, на массовые акции почти не реагируют",
        "customers": 17_000, "freq": 0.03, "aov": 12_000,
        "budget_before": 6.0, "budget_after": 1.0,
        "freq_after": 0.025, "aov_after": 12_000, "extra_revenue": 0.0,
        "goal": "Перестать сжигать бюджет: одна попытка, затем исключение",
        "lever": "Одна недорогая реактивация, после неё — исключение из "
                 "платного ретаргета и массовых промокодов",
        "channels": "Один e-mail с лучшим предложением, далее только "
                    "бесплатные каналы",
        "extra_note": "",
    },
]

ACQUISITION = {
    "id": "acquisition", "name": "Привлечение новых",
    "budget_before": 12.5, "budget_after": 11.0,
    "lever": "Исключить текущих клиентов из платных аудиторий, "
             "look-alike на чемпионах, отключить кампании с CAC выше LTV",
}


def segment_calc(seg: dict, share: float = 1.0) -> dict:
    revenue_before = seg["customers"] * seg["freq"] * seg["aov"] / 1e6
    freq = seg["freq"] + (seg["freq_after"] - seg["freq"]) * share
    aov = seg["aov"] + (seg["aov_after"] - seg["aov"]) * share
    revenue_after = seg["customers"] * freq * aov / 1e6 + seg["extra_revenue"] * share
    return {
        "revenue_before": revenue_before,
        "revenue_after": revenue_after,
        "delta": revenue_after - revenue_before,
    }


def build() -> dict:
    new_revenue = NEW_CUSTOMERS * NEW_AOV / 1e6
    budget_before = sum(s["budget_before"] for s in SEGMENTS) + ACQUISITION["budget_before"]
    budget_after = sum(s["budget_after"] for s in SEGMENTS) + ACQUISITION["budget_after"]
    total_customers = sum(s["customers"] for s in SEGMENTS)

    segments = []
    for seg in SEGMENTS:
        calc = segment_calc(seg)
        segments.append({
            **seg, **calc,
            "share_customers": seg["customers"] / total_customers,
            "budget_delta": seg["budget_after"] - seg["budget_before"],
        })

    revenue_before = sum(s["revenue_before"] for s in segments) + new_revenue
    for s in segments:
        s["share_revenue"] = s["revenue_before"] / revenue_before

    scenarios = {}
    for name, share in SCENARIO_SHARE.items():
        after = sum(segment_calc(s, share)["revenue_after"] for s in SEGMENTS) + new_revenue
        delta = after - revenue_before
        scenarios[name] = {
            "share_of_uplift": share,
            "revenue_after": after,
            "revenue_growth": after / revenue_before - 1,
            "margin_effect": delta * GROSS_MARGIN + (budget_before - budget_after),
        }

    base = scenarios["base"]
    return {
        "assumptions": {
            "customers_12m": total_customers,
            "gross_margin": GROSS_MARGIN,
            "new_customers_month": NEW_CUSTOMERS,
            "new_aov": NEW_AOV,
            "new_revenue": new_revenue,
        },
        "segments": segments,
        "acquisition": {**ACQUISITION,
                        "budget_delta": ACQUISITION["budget_after"] - ACQUISITION["budget_before"]},
        "totals": {
            "budget_before": budget_before,
            "budget_after": budget_after,
            "budget_change": budget_after / budget_before - 1,
            "revenue_before": revenue_before,
            "revenue_after": base["revenue_after"],
            "revenue_growth": base["revenue_growth"],
            "crr_before": budget_before / revenue_before,
            "crr_after": budget_after / base["revenue_after"],
            "romi_before": revenue_before / budget_before,
            "romi_after": base["revenue_after"] / budget_after,
            "margin_effect": base["margin_effect"],
        },
        "scenarios": scenarios,
    }


def write_csv(model: dict) -> None:
    cols = ["name", "rule", "customers", "freq", "aov", "revenue_before",
            "budget_before", "budget_after", "freq_after", "aov_after",
            "extra_revenue", "revenue_after", "delta"]
    with open(OUT / "segments.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for s in model["segments"]:
            w.writerow([round(s[c], 4) if isinstance(s[c], float) else s[c] for c in cols])


def write_xlsx(model: dict) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Модель"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="006A4E")
    soft = PatternFill("solid", fgColor="DAFFED")

    ws["A1"] = "RFM-перераспределение бюджета · месячная модель, млн ₸"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Мятные ячейки — допущения, меняйте их: всё остальное пересчитается."

    headers = ["Сегмент", "RFM-правило", "Клиенты", "Частота до", "Чек до, ₸",
               "Выручка до", "Бюджет до", "Бюджет после", "Частота после",
               "Чек после, ₸", "Доп. выручка", "Выручка после", "Δ выручки",
               "Δ бюджета", "Выручка на 1 ₸ бюджета после"]
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font, c.fill = head, fill
        c.alignment = Alignment(wrap_text=True, vertical="top")

    first = 5
    for r, s in enumerate(model["segments"], first):
        vals = [s["name"], s["rule"], s["customers"], s["freq"], s["aov"], None,
                s["budget_before"], s["budget_after"], s["freq_after"],
                s["aov_after"], s["extra_revenue"]]
        for col, v in enumerate(vals, 1):
            if v is not None:
                ws.cell(row=r, column=col, value=v)
        for col in (3, 4, 5, 7, 8, 9, 10, 11):
            ws.cell(row=r, column=col).fill = soft
        ws.cell(row=r, column=6, value=f"=C{r}*D{r}*E{r}/1000000")
        ws.cell(row=r, column=12, value=f"=C{r}*I{r}*J{r}/1000000+K{r}")
        ws.cell(row=r, column=13, value=f"=L{r}-F{r}")
        ws.cell(row=r, column=14, value=f"=H{r}-G{r}")
        ws.cell(row=r, column=15, value=f"=L{r}/H{r}")

    acq = first + len(model["segments"])
    ws.cell(row=acq, column=1, value="Новые клиенты (привлечение)")
    ws.cell(row=acq, column=2, value="—")
    ws.cell(row=acq, column=3, value=model["assumptions"]["new_customers_month"]).fill = soft
    ws.cell(row=acq, column=4, value=1).fill = soft
    ws.cell(row=acq, column=5, value=model["assumptions"]["new_aov"]).fill = soft
    ws.cell(row=acq, column=6, value=f"=C{acq}*D{acq}*E{acq}/1000000")
    ws.cell(row=acq, column=7, value=model["acquisition"]["budget_before"]).fill = soft
    ws.cell(row=acq, column=8, value=model["acquisition"]["budget_after"]).fill = soft
    ws.cell(row=acq, column=9, value=f"=D{acq}")
    ws.cell(row=acq, column=10, value=f"=E{acq}")
    ws.cell(row=acq, column=11, value=0)
    ws.cell(row=acq, column=12, value=f"=C{acq}*I{acq}*J{acq}/1000000+K{acq}")
    ws.cell(row=acq, column=13, value=f"=L{acq}-F{acq}")
    ws.cell(row=acq, column=14, value=f"=H{acq}-G{acq}")
    ws.cell(row=acq, column=15, value=f"=L{acq}/H{acq}")

    tot = acq + 1
    ws.cell(row=tot, column=1, value="Итого").font = Font(bold=True)
    for col in "CFGHKLMN":
        ws[f"{col}{tot}"] = f"=SUM({col}{first}:{col}{acq})"
        ws[f"{col}{tot}"].font = Font(bold=True)

    k = tot + 2
    rows = [
        ("Изменение бюджета", f"=H{tot}/G{tot}-1", "0.0%"),
        ("Рост выручки", f"=L{tot}/F{tot}-1", "0.0%"),
        ("Доля маркетинга в выручке до (CRR)", f"=G{tot}/F{tot}", "0.0%"),
        ("Доля маркетинга в выручке после (CRR)", f"=H{tot}/L{tot}", "0.0%"),
        ("Валовая маржа", model["assumptions"]["gross_margin"], "0%"),
        ("Эффект на валовую прибыль, млн ₸/мес", f"=M{tot}*B{k + 4}+(G{tot}-H{tot})", "0.0"),
        ("Цель: бюджет −20% выполнена", f'=IF(B{k}<=-0.2,"да","нет")', None),
        ("Цель: выручка +10% выполнена", f'=IF(B{k + 1}>=0.1,"да","нет")', None),
    ]
    for i, (label, formula, fmt) in enumerate(rows):
        ws.cell(row=k + i, column=1, value=label)
        c = ws.cell(row=k + i, column=2, value=formula)
        if fmt:
            c.number_format = fmt
    ws.cell(row=k + 4, column=2).fill = soft

    for row in ws.iter_rows(min_row=first, max_row=tot, min_col=6, max_col=15):
        for c in row:
            c.number_format = "0.00"
    widths = [26, 14, 10, 10, 10, 11, 10, 11, 11, 11, 11, 12, 11, 11, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    sc = wb.create_sheet("Сценарии")
    sc.append(["Сценарий", "Доля реализации целевого эффекта", "Выручка после, млн ₸",
               "Рост выручки", "Эффект на валовую прибыль, млн ₸"])
    labels = {"pessimistic": "Пессимистичный", "base": "Базовый", "optimistic": "Оптимистичный"}
    for key, v in model["scenarios"].items():
        sc.append([labels[key], v["share_of_uplift"], round(v["revenue_after"], 2),
                   round(v["revenue_growth"], 4), round(v["margin_effect"], 2)])
    for col, w in zip("ABCDE", (18, 18, 18, 14, 22)):
        sc.column_dimensions[col].width = w

    wb.save(OUT / "rfm_budget_model.xlsx")
    return True


def main() -> dict:
    OUT.mkdir(exist_ok=True)
    model = build()
    (OUT / "model.json").write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(model)
    xlsx = write_xlsx(model)
    t = model["totals"]
    print(f"budget {t['budget_before']:.1f} -> {t['budget_after']:.1f} ({t['budget_change']:+.1%})")
    print(f"revenue {t['revenue_before']:.2f} -> {t['revenue_after']:.2f} ({t['revenue_growth']:+.1%})")
    for k, v in model["scenarios"].items():
        print(f"  {k}: {v['revenue_growth']:+.1%}, margin effect {v['margin_effect']:.1f}")
    print("xlsx written" if xlsx else "xlsx skipped (pip install openpyxl)")
    return model


if __name__ == "__main__":
    main()
